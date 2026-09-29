"""
Cloud Synchronization Module
GOOGLE SHEETS VERSION - COMPLETE
- Records searches to Google Sheets (Risk Searches sheet)
- Reads keywords from Google Sheets (Keywords sheet)
- Maintains local Excel backup for user's searches
- Auto-syncs keyword changes from Google Sheets
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import gspread
from google.oauth2.service_account import Credentials
from datetime import datetime
import json
import time
import getpass
from config import GOOGLE_SHEETS_CONFIG, LOCAL_FILES, EXCEL_HEADERS

RUSSIAN_KEYWORDS_WORKSHEET_NAME = GOOGLE_SHEETS_CONFIG.get(
    'russian_keywords_worksheet_name',
    'Russian_Keyword_Sheet'
)


class CloudSync:
    """Manages synchronization with Google Sheets"""
    
    def __init__(self):
        self.username = getpass.getuser()
        self.client = None
        self.sheet = None
        self.worksheet = None
        self.keywords_worksheet = None
        self.russian_keywords_worksheet = None
        self._connect_to_google_sheets()
    
    def _connect_to_google_sheets(self):
        """Connect to Google Sheets and setup worksheets"""
        try:
            credentials_file = GOOGLE_SHEETS_CONFIG.get('credentials_file')
            sheet_id = GOOGLE_SHEETS_CONFIG.get('sheet_id')
            
            if not credentials_file or not os.path.exists(credentials_file):
                print("⚠️ Google Sheets credentials not found - using local mode")
                return
            
            if not sheet_id or sheet_id == "YOUR_GOOGLE_SHEET_ID":
                print("⚠️ Google Sheet ID not configured - using local mode")
                return
            
            # Define the scopes
            scopes = [
                'https://www.googleapis.com/auth/spreadsheets',
                'https://www.googleapis.com/auth/drive'
            ]
            
            # Load credentials
            creds = Credentials.from_service_account_file(credentials_file, scopes=scopes)
            
            # Authorize the client
            self.client = gspread.authorize(creds)
            
            # Open the spreadsheet by key (ID)
            self.sheet = self.client.open_by_key(sheet_id)
            
            # Setup Risk Searches worksheet
            try:
                self.worksheet = self.sheet.worksheet("Risk Searches")
            except gspread.exceptions.WorksheetNotFound:
                # Create worksheet if doesn't exist
                self.worksheet = self.sheet.add_worksheet(title="Risk Searches", rows=1000, cols=10)
                # Add headers
                headers = ['ID', 'Client Name', 'Report Date', 'Username', 
                          'Search URL', 'Status', 'Keywords Used', 'App Version']
                self.worksheet.update('A1:H1', [headers])
                # Format header row
                self.worksheet.format('A1:H1', {
                    'backgroundColor': {'red': 0.27, 'green': 0.45, 'blue': 0.77},
                    'textFormat': {'bold': True, 'foregroundColor': {'red': 1, 'green': 1, 'blue': 1}}
                })
                print("📄 Created 'Risk Searches' worksheet")
            
            # Setup Keywords worksheet
            try:
                self.keywords_worksheet = self.sheet.worksheet("Keywords")
                print("✅ Found 'Keywords' worksheet")
            except gspread.exceptions.WorksheetNotFound:
                # Create Keywords worksheet
                self.keywords_worksheet = self.sheet.add_worksheet(title="Keywords", rows=100, cols=3)
                # Add headers
                headers = ['Category', 'Keyword', 'Active']
                self.keywords_worksheet.update('A1:C1', [headers])
                
                # Format header row
                self.keywords_worksheet.format('A1:C1', {
                    'backgroundColor': {'red': 0.27, 'green': 0.45, 'blue': 0.77},
                    'textFormat': {'bold': True, 'foregroundColor': {'red': 1, 'green': 1, 'blue': 1}}
                })
                
                # Add default 32 keywords
                self._populate_default_keywords()
                print("📄 Created 'Keywords' worksheet with 32 default keywords")

            try:
                self.russian_keywords_worksheet = self.sheet.worksheet(
                    RUSSIAN_KEYWORDS_WORKSHEET_NAME
                )
                russian_sheet_values = self.russian_keywords_worksheet.get_all_values()
                if not russian_sheet_values:
                    self._initialize_russian_keywords_worksheet()
                elif (
                    len(russian_sheet_values) == 1
                    and [cell.strip().lower() for cell in russian_sheet_values[0][:3]]
                    == ['category', 'keyword', 'active']
                ):
                    self._populate_russian_keywords()
                print("Found 'Russian_Keyword_Sheet' worksheet")
            except gspread.exceptions.WorksheetNotFound:
                self.russian_keywords_worksheet = self.sheet.add_worksheet(
                    title=RUSSIAN_KEYWORDS_WORKSHEET_NAME,
                    rows=100,
                    cols=3
                )
                self._initialize_russian_keywords_worksheet()
                print("Created 'Russian_Keyword_Sheet' with 20 keywords")
            except Exception as e:
                self.russian_keywords_worksheet = None
                print(f"❌ Russian keyword sheet setup failed: {str(e)}")
                print("⚠️ The standard Keywords worksheet remains available")
            
            print("✅ Connected to Google Sheets")
            
        except Exception as e:
            print(f"❌ Google Sheets connection failed: {str(e)}")
            print(f"   Error details: {type(e).__name__}")
            print("⚠️ Running in local mode")
            self.worksheet = None
            self.keywords_worksheet = None
    
    def _populate_default_keywords(self):
        """Populate Keywords worksheet with default 32 keywords"""
        
        keywords_data = [
            ['Financial Crime', 'fraud', 'TRUE'],
            ['Financial Crime', 'corruption', 'TRUE'],
            ['Financial Crime', 'bribery', 'TRUE'],
            ['Financial Crime', 'money laundering', 'TRUE'],
            ['Financial Crime', 'terrorist financing', 'TRUE'],
            ['Financial Crime', 'sanctions violation', 'TRUE'],
            ['Financial Crime', 'embezzlement', 'TRUE'],
            ['Financial Crime', 'tax evasion', 'TRUE'],
            ['Legal & Regulatory', 'lawsuit', 'TRUE'],
            ['Legal & Regulatory', 'litigation', 'TRUE'],
            ['Legal & Regulatory', 'court case', 'TRUE'],
            ['Legal & Regulatory', 'regulatory action', 'TRUE'],
            ['Legal & Regulatory', 'enforcement action', 'TRUE'],
            ['Legal & Regulatory', 'compliance breach', 'TRUE'],
            ['Human Rights & Labour', 'human rights violation', 'TRUE'],
            ['Human Rights & Labour', 'forced labour', 'TRUE'],
            ['Human Rights & Labour', 'child labour', 'TRUE'],
            ['Human Rights & Labour', 'discrimination at work', 'TRUE'],
            ['Human Rights & Labour', 'labour law violation', 'TRUE'],
            ['Human Rights & Labour', 'union suppression', 'TRUE'],
            ['Human Rights & Labour', 'collective bargaining restriction', 'TRUE'],
            ['Human Rights & Labour', 'unsafe working conditions', 'TRUE'],
            ['Environmental & ESG', 'environmental damage', 'TRUE'],
            ['Environmental & ESG', 'pollution incident', 'TRUE'],
            ['Environmental & ESG', 'environmental negligence', 'TRUE'],
            ['Environmental & ESG', 'toxic waste', 'TRUE'],
            ['Governance & Ethics', 'management misconduct', 'TRUE'],
            ['Governance & Ethics', 'governance failure', 'TRUE'],
            ['Governance & Ethics', 'ethics violation', 'TRUE'],
            ['Governance & Ethics', 'whistleblower allegation', 'TRUE'],
            ['Cyber & Data', 'data breach', 'TRUE'],
            ['Cyber & Data', 'leaked documents', 'TRUE'],
        ]
        
        # Append all keywords at once (starts from row 2)
        self.keywords_worksheet.append_rows(keywords_data)
        
        print("✅ Added 32 default keywords to Google Sheet")
    
    def _populate_russian_keywords(self):
        """Populate the Russian risk keyword worksheet with its default terms"""
        keywords_data = [
            ['Russian Risk', 'Blacklist', 'TRUE'],
            ['Russian Risk', 'Breach', 'TRUE'],
            ['Russian Risk', 'Bribery', 'TRUE'],
            ['Russian Risk', 'Crime', 'TRUE'],
            ['Russian Risk', 'Criminal', 'TRUE'],
            ['Russian Risk', 'Controversy', 'TRUE'],
            ['Russian Risk', 'Controversial', 'TRUE'],
            ['Russian Risk', 'Dispute', 'TRUE'],
            ['Russian Risk', 'Drug', 'TRUE'],
            ['Russian Risk', 'Fraud', 'TRUE'],
            ['Russian Risk', 'Hacking', 'TRUE'],
            ['Russian Risk', 'Litigation', 'TRUE'],
            ['Russian Risk', 'Money Laundering', 'TRUE'],
            ['Russian Risk', 'Proliferation', 'TRUE'],
            ['Russian Risk', 'Smuggling', 'TRUE'],
            ['Russian Risk', 'Terrorism', 'TRUE'],
            ['Russian Risk', 'Terrorist', 'TRUE'],
            ['Russian Risk', 'Trafficking', 'TRUE'],
            ['Russian Risk', 'Weapon of Mass Destruction', 'TRUE'],
            ['Russian Risk', 'Russia', 'TRUE'],
        ]
        self.russian_keywords_worksheet.append_rows(keywords_data)

    def _initialize_russian_keywords_worksheet(self):
        """Set up headers and default keywords on a new or empty worksheet"""
        self.russian_keywords_worksheet.update(
            'A1:C1',
            [['Category', 'Keyword', 'Active']]
        )
        self.russian_keywords_worksheet.format('A1:C1', {
            'backgroundColor': {'red': 0.27, 'green': 0.45, 'blue': 0.77},
            'textFormat': {
                'bold': True,
                'foregroundColor': {'red': 1, 'green': 1, 'blue': 1}
            }
        })
        self._populate_russian_keywords()

    def get_keyword_sheets(self):
        """Return keyword worksheets that users can select in the search form"""
        return [
            {
                'name': GOOGLE_SHEETS_CONFIG['keywords_worksheet_name'],
                'label': 'Standard Keywords'
            },
            {
                'name': RUSSIAN_KEYWORDS_WORKSHEET_NAME,
                'label': 'Russian_Keyword_Sheet'
            }
        ]

    @staticmethod
    def _extract_active_keywords(all_values):
        keywords = []
        for row in all_values[1:]:
            if len(row) >= 3:
                keyword = row[1]
                active = row[2].upper()
                if active in ['TRUE', 'YES', '1', 'CHECKED', 'X']:
                    if keyword and keyword.strip():
                        keywords.append(keyword.strip())
        return keywords

    def get_keywords(self, worksheet_name=None):
        """
        Get active keywords from the selected Google Sheets worksheet
        Reads all active keywords (where Active column = TRUE)
        """

        standard_sheet_name = GOOGLE_SHEETS_CONFIG['keywords_worksheet_name']
        russian_sheet_name = RUSSIAN_KEYWORDS_WORKSHEET_NAME
        worksheet_name = worksheet_name or standard_sheet_name

        if worksheet_name == russian_sheet_name:
            if self.russian_keywords_worksheet is None:
                raise RuntimeError(
                    f"'{russian_sheet_name}' is unavailable. Check the Google Sheets connection."
                )
            try:
                keywords = self._extract_active_keywords(
                    self.russian_keywords_worksheet.get_all_values()
                )
                print(f"Loaded {len(keywords)} keywords from '{russian_sheet_name}'")
                return keywords
            except Exception as e:
                raise RuntimeError(
                    f"Could not read keywords from '{russian_sheet_name}': {str(e)}"
                ) from e

        if worksheet_name != standard_sheet_name:
            raise ValueError(f"Unknown keyword worksheet: {worksheet_name}")

        if self.keywords_worksheet is not None:
            try:
                # Get all values from Keywords sheet
                all_values = self.keywords_worksheet.get_all_values()
                keywords = self._extract_active_keywords(all_values)
                print(f"📋 Loaded {len(keywords)} keywords from Google Sheets")
                return keywords
            
            except Exception as e:
                print(f"❌ Error reading keywords from Google Sheets: {str(e)}")
                print("⚠️ Using default keywords")
                return self._get_default_keywords()
        else:
            # No Google Sheets connection - use default
            print("⚠️ Using default keywords (no Google Sheets)")
            return self._get_default_keywords()
    
    def _get_default_keywords(self):
        """Fallback keywords if Google Sheets unavailable"""
        return [
            "Bankruptcy","Insolvency","Liquidation","Winding-up petition","Receivership","Crime"
            ,"Fraud","Corruption","Bribery","Money laundering","Terrorist financing","Sanctions violation"
            ,"Embezzlement","Tax evasion","Human rights abuse","Labour rights violation","Forced labour"
            ,"Child labour","Workplace discrimination","Union suppression","Environmental violation","Misconduct"
            ,"Governance failure","Ethics violation","Whistleblower allegation","Reputation risk","Data breach"
            ,"Leaked documents","Lawsuit","Court case","Litigation","Regulatory","Enforcement action"
        ]
    
    def append_search_record(self, record_data):
        """
        Append search record to Google Sheets AND local Excel
        Returns record ID
        """
        
        # Always save to local Excel first
        local_id = self._save_to_local_excel(record_data)
        
        # Try to save to Google Sheets
        if self.worksheet is not None:
            try:
                # Get current row count
                all_values = self.worksheet.get_all_values()
                next_row = len(all_values) + 1
                record_id = next_row - 1  # Subtract header row
                
                # Prepare row data
                row_data = [
                    record_id,
                    record_data.get('client_name', ''),
                    record_data.get('report_date', ''),
                    self.username,
                    record_data.get('search_url', ''),
                    record_data.get('status', 'Success'),
                    record_data.get('keywords_used', 0),
                    record_data.get('app_version', '')
                ]
                
                # Append row to Google Sheets
                self.worksheet.append_row(row_data, value_input_option='USER_ENTERED')
                
                print(f"✅ Saved to Google Sheets (Row {next_row}, ID: {record_id})")
                print(f"💾 Also saved to local Excel backup")
                
                return record_id
            
            except Exception as e:
                print(f"❌ Google Sheets save failed: {str(e)}")
                print(f"💾 Data saved to local Excel only (ID: {local_id})")
                return local_id
        else:
            # No Google Sheets connection
            print(f"💾 Saved to local Excel only (ID: {local_id})")
            return local_id
    
    def _save_to_local_excel(self, record_data):
        """Save search to local Excel file as backup"""
        
        # Get local file path for this user
        cache_dir = os.path.join(os.path.expanduser("~"), ".risk_search_tool")
        os.makedirs(cache_dir, exist_ok=True)
        local_file = os.path.join(cache_dir, f"{self.username}_searches.xlsx")
        
        # Create file if doesn't exist
        if not os.path.exists(local_file):
            self._create_local_excel_file(local_file)
        
        try:
            wb = openpyxl.load_workbook(local_file)
            ws = wb.active
            
            next_row = ws.max_row + 1
            next_id = next_row - 1
            
            # Add data
            ws.cell(row=next_row, column=1).value = next_id
            ws.cell(row=next_row, column=2).value = record_data.get('client_name')
            ws.cell(row=next_row, column=3).value = record_data.get('report_date')
            ws.cell(row=next_row, column=4).value = self.username
            ws.cell(row=next_row, column=5).value = record_data.get('search_url')
            ws.cell(row=next_row, column=6).value = record_data.get('status')
            ws.cell(row=next_row, column=7).value = record_data.get('keywords_used')
            ws.cell(row=next_row, column=8).value = record_data.get('app_version')
            
            wb.save(local_file)
            wb.close()
            
            return next_id
        
        except Exception as e:
            print(f"❌ Local save failed: {str(e)}")
            return None
    
    def _create_local_excel_file(self, filepath):
        """Create new local Excel file with headers"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{self.username} Searches"
        
        # Add headers with styling
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        
        for col_num, header in enumerate(EXCEL_HEADERS['master_records'], 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
        
        # Set column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 22
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 70
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 15
        
        wb.save(filepath)
        wb.close()
        print(f"📄 Created local Excel backup: {filepath}")
    
    def get_all_searches(self):
        """Get ALL searches from Google Sheets"""
        
        if self.worksheet is not None:
            try:
                # Get all values from sheet
                all_values = self.worksheet.get_all_values()
                
                # Skip header row
                searches = []
                for row in all_values[1:]:
                    if len(row) >= 8 and row[1]:  # Has data in client name column
                        searches.append({
                            'id': row[0] if row[0] else '',
                            'client_name': row[1],
                            'report_date': row[2],
                            'username': row[3],
                            'search_url': row[4],
                            'status': row[5],
                            'keywords_used': row[6],
                            'app_version': row[7]
                        })
                
                return searches
            
            except Exception as e:
                print(f"❌ Error loading from Google Sheets: {str(e)}")
                return self._get_local_searches()
        else:
            return self._get_local_searches()
    
    def _get_local_searches(self):
        """Get searches from local Excel file"""
        
        cache_dir = os.path.join(os.path.expanduser("~"), ".risk_search_tool")
        local_file = os.path.join(cache_dir, f"{self.username}_searches.xlsx")
        
        if not os.path.exists(local_file):
            return []
        
        try:
            wb = openpyxl.load_workbook(local_file, read_only=True, data_only=True)
            ws = wb.active
            
            searches = []
            for row in range(2, ws.max_row + 1):
                searches.append({
                    'id': ws.cell(row=row, column=1).value,
                    'client_name': ws.cell(row=row, column=2).value,
                    'report_date': ws.cell(row=row, column=3).value,
                    'username': ws.cell(row=row, column=4).value,
                    'search_url': ws.cell(row=row, column=5).value,
                    'status': ws.cell(row=row, column=6).value,
                    'keywords_used': ws.cell(row=row, column=7).value,
                    'app_version': ws.cell(row=row, column=8).value
                })
            
            wb.close()
            return searches
        
        except Exception as e:
            return []
    
    def get_recent_searches(self, limit=10):
        """Get recent searches"""
        
        all_searches = self.get_all_searches()
        return all_searches[-limit:] if len(all_searches) > limit else all_searches
    
    def get_total_records(self):
        """Get total number of records"""
        
        if self.worksheet is not None:
            try:
                all_values = self.worksheet.get_all_values()
                return len(all_values) - 1  # Subtract header
            except Exception as e:
                return self._get_local_record_count()
        else:
            return self._get_local_record_count()
    
    def _get_local_record_count(self):
        """Count records in local Excel"""
        cache_dir = os.path.join(os.path.expanduser("~"), ".risk_search_tool")
        local_file = os.path.join(cache_dir, f"{self.username}_searches.xlsx")
        
        if os.path.exists(local_file):
            try:
                wb = openpyxl.load_workbook(local_file, read_only=True)
                ws = wb.active
                total = ws.max_row - 1
                wb.close()
                return total
            except:
                return 0
        return 0
    
    def export_to_excel(self, filepath, username_filter=None):
        """Export data to Excel file"""
        
        # Create new workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        
        if username_filter:
            ws.title = f"{username_filter} Searches"
        else:
            ws.title = "All Searches"
        
        # Style settings
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Add headers
        headers = ['ID', 'Client Name', 'Report Date', 'Username', 'Search URL', 
                  'Status', 'Keywords Used', 'App Version']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Get data
        all_searches = self.get_all_searches()
        
        # Filter by username if specified
        if username_filter:
            all_searches = [s for s in all_searches if s.get('username') == username_filter]
        
        # Add data rows
        for row_num, search in enumerate(all_searches, 2):
            ws.cell(row=row_num, column=1).value = search.get('id')
            ws.cell(row=row_num, column=1).border = border
            
            ws.cell(row=row_num, column=2).value = search.get('client_name')
            ws.cell(row=row_num, column=2).border = border
            
            ws.cell(row=row_num, column=3).value = search.get('report_date')
            ws.cell(row=row_num, column=3).border = border
            
            ws.cell(row=row_num, column=4).value = search.get('username')
            ws.cell(row=row_num, column=4).border = border
            
            search_url = search.get('search_url')
            ws.cell(row=row_num, column=5).value = search_url
            if search_url:
                ws.cell(row=row_num, column=5).hyperlink = search_url
                ws.cell(row=row_num, column=5).font = Font(color="0563C1", underline="single")
            ws.cell(row=row_num, column=5).border = border
            
            ws.cell(row=row_num, column=6).value = search.get('status')
            ws.cell(row=row_num, column=6).border = border
            
            ws.cell(row=row_num, column=7).value = search.get('keywords_used')
            ws.cell(row=row_num, column=7).border = border
            
            ws.cell(row=row_num, column=8).value = search.get('app_version')
            ws.cell(row=row_num, column=8).border = border
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 22
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 70
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 15
        
        # Save file
        wb.save(filepath)
        wb.close()
        
        print(f"✅ Excel exported to: {filepath}")
        return filepath


if __name__ == "__main__":
    print("Testing Google Sheets connection and keyword reading...\n")
    
    cloud = CloudSync()
    
    if cloud.keywords_worksheet:
        keywords = cloud.get_keywords()
        print(f"\n✅ Successfully loaded {len(keywords)} keywords:")
        for i, kw in enumerate(keywords, 1):
            print(f"  {i}. {kw}")
    else:
        print("\n⚠️ Keywords worksheet not available")